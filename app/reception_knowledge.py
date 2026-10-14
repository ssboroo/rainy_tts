"""Per-organisation receptionist knowledge and safe document connectors.

Knowledge is treated strictly as untrusted *data*, not agent instructions.
Return short verbatim passages with source names; never fabricate answers.
Only admin controls documents; existing agent bearer-URL tools can search them.
"""
import asyncio
import csv
import io
import json
import os
import re
from zipfile import ZipFile,BadZipFile
import time
from datetime import datetime, timedelta
from pathlib import Path
from urllib.parse import urlsplit, parse_qs

import httpx
from fastapi import HTTPException, Request, UploadFile, File
from . import core

MAX_BYTES=5*1024*1024
MAX_TEXT=180_000
MAX_CHUNKS=350
SHEET_INTERVAL=6*3600
EXTS={".txt",".md",".csv",".xlsx",".docx",".pdf"}
SHEET_PATH=re.compile(r"^/spreadsheets/d/([A-Za-z0-9_-]{12,160})(?:/(?:edit|view|export|pubhtml))?/?$")
STOP={"нь","бол","юу","яаж","хэд","ямар","вэ","байна","талаар","хаана","чадах","тухай","манай","танай","гэж","хэн","хэнтэй","байгаа","болох","хэдэн","тэгэх","энэ","тэр","does","how","the","is","what","where","a","do","we","with"}
def ensure_schema():
    with core.db() as db:
        db.executescript("""
        CREATE TABLE IF NOT EXISTS reception_business(
          user_id TEXT PRIMARY KEY REFERENCES users(id) ON DELETE CASCADE,
          name TEXT NOT NULL DEFAULT '', address TEXT NOT NULL DEFAULT '',
          contact_phone TEXT NOT NULL DEFAULT '', intro TEXT NOT NULL DEFAULT '',
          open_hour INTEGER NOT NULL DEFAULT 9, close_hour INTEGER NOT NULL DEFAULT 18,
          closed_weekdays TEXT NOT NULL DEFAULT '[6]',
          updated REAL NOT NULL
        );
        CREATE TABLE IF NOT EXISTS reception_sources(
          id TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE,
          name TEXT NOT NULL, kind TEXT NOT NULL, sheet_url TEXT, created REAL NOT NULL,
          refreshed REAL NOT NULL, last_error TEXT NOT NULL DEFAULT ''
        );
        CREATE INDEX IF NOT EXISTS reception_sources_user ON reception_sources(user_id,created DESC);
        CREATE TABLE IF NOT EXISTS reception_passages(
          source_id TEXT NOT NULL REFERENCES reception_sources(id) ON DELETE CASCADE,
          line_no INTEGER NOT NULL, text TEXT NOT NULL, PRIMARY KEY(source_id,line_no)
        );
        """)

def read_business(user_id):
    ensure_schema()
    with core.db() as db:
        row=db.execute("SELECT name,address,contact_phone,intro,open_hour,close_hour,closed_weekdays FROM reception_business WHERE user_id=?",(user_id,)).fetchone()
    if not row:
        return {"name":"","address":"","contact_phone":"","intro":"","open_hour":9,"close_hour":18,"closed_weekdays":[6]}
    result=dict(row)
    result["closed_weekdays"]=json.loads(result["closed_weekdays"])
    return result

def save_business(user_id,payload):
    if not isinstance(payload,dict):raise HTTPException(422,"Байгууллагын мэдээлэл буруу байна.")
    values={}
    for key,limit in [("name",120),("address",500),("contact_phone",30),("intro",1000)]:
        val=payload.get(key,"")
        if not isinstance(val,str) or len(val)>limit:
            raise HTTPException(422,"Байгууллагын мэдээлэл хэт урт эсвэл буруу байна.")
        values[key]=val.strip()
    if not values["name"]:raise HTTPException(422,"Байгууллагын нэр шаардлагатай.")
    opening=payload.get("open_hour",9)
    closing=payload.get("close_hour",18)
    closed=payload.get("closed_weekdays",[6])
    if (type(opening)!=int or type(closing)!=int or not (0<=opening<=22 and opening+1<=closing<=23)
        or not isinstance(closed,list) or len(closed)>7 or any(type(x)!=int or x not in range(7) for x in closed)
        or len(set(closed))!=len(closed)):
        raise HTTPException(422,"Ажиллах цаг 0–23, амралтын гарагууд 0–6 (Даваа=0) байна.")
    values.update({"open_hour":opening,"close_hour":closing,"closed_weekdays":closed})
    with core.db() as db:
        db.execute("""INSERT INTO reception_business(user_id,name,address,contact_phone,intro,open_hour,close_hour,closed_weekdays,updated)
                      VALUES(?,?,?,?,?,?,?,?,?)
                      ON CONFLICT(user_id) DO UPDATE SET
                      name=excluded.name,address=excluded.address,contact_phone=excluded.contact_phone,
                      intro=excluded.intro,open_hour=excluded.open_hour,close_hour=excluded.close_hour,
                      closed_weekdays=excluded.closed_weekdays,updated=excluded.updated""",
                   (user_id,values["name"],values["address"],values["contact_phone"],values["intro"],
                    opening,closing,json.dumps(closed),time.time()))
    return values

def extract_document(filename,data):
    suffix=Path(filename).suffix.lower()
    if suffix not in EXTS:
        raise HTTPException(422,"TXT, CSV, XLSX, DOCX, PDF дэмжинэ. Хуучин DOC/XLS-г шинэ форматаар хадгална уу.")
    if not data or len(data)>MAX_BYTES:
        raise HTTPException(413 if len(data)>MAX_BYTES else 422,"Файлын хэмжээ 5 MB-аас их эсвэл хоосон байна.")
    # Reject compressed XML bombs before DOCX/XLSX parsers expand archives.
    if suffix in (".xlsx",".docx"):
        try:
            with ZipFile(io.BytesIO(data)) as archive:
                infos=archive.infolist()
                if len(infos)>5000 or sum(item.file_size for item in infos)>20*1024*1024:
                    raise HTTPException(413,"Баримтын задлагдах хэмжээ хэтэрсэн байна.")
        except HTTPException:raise
        except BadZipFile as exc:
            raise HTTPException(422,"DOCX/XLSX архив гэмтсэн байна.") from exc
    try:
        if suffix in (".txt",".md"):
            lines=data.decode("utf-8-sig").splitlines()
        elif suffix==".csv":
            rows=list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
            lines=[" | ".join(str(cell).strip() for cell in row[:24] if str(cell).strip()) for row in rows[:5000]]
        elif suffix==".xlsx":
            from openpyxl import load_workbook
            book=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
            lines=[]
            try:
                for sheet in book.worksheets[:12]:
                    lines.append("Хүснэгтийн хуудас: "+sheet.title)
                    for row in sheet.iter_rows(max_row=2500,max_col=24,values_only=True):
                        segment=" | ".join(str(v).strip()[:400] for v in row if v is not None and str(v).strip())
                        if segment:lines.append(segment)
                        if sum(map(len,lines))>MAX_TEXT:break
            finally:book.close()
        elif suffix==".docx":
            from docx import Document
            doc=Document(io.BytesIO(data))
            lines=[x.text.strip() for x in doc.paragraphs if x.text.strip()]
            for table in doc.tables[:30]:
                for row in table.rows[:500]:
                    lines.append(" | ".join(cell.text.strip()[:400] for cell in row.cells if cell.text.strip()))
        elif suffix==".pdf":
            from pypdf import PdfReader
            reader=PdfReader(io.BytesIO(data),strict=False)
            if reader.is_encrypted:
                raise HTTPException(422,"Нууц үгтэй PDF уншихгүй.")
            if len(reader.pages)>100:
                raise HTTPException(422,"PDF хамгийн ихдээ 100 хуудастай байна.")
            lines=[]
            for page in reader.pages:
                text=page.extract_text() or ""
                lines.extend(text.splitlines()[:1500])
                if sum(map(len,lines))>MAX_TEXT:break
        else:lines=[]
    except HTTPException:raise
    except (UnicodeError,ValueError,KeyError,TypeError,RuntimeError,Exception) as exc:
        # Includes parser errors; do not disclose paths or embedded content.
        raise HTTPException(422,"Файлын агуулгыг уншиж чадсангүй. Зөв формат, гэмтээгүй файл оруулна уу.") from exc
    cleaned=[]
    total=0
    for raw in lines:
        s=" ".join(str(raw).split())[:900]
        if not s:continue
        total+=len(s)
        if total>MAX_TEXT or len(cleaned)>=MAX_CHUNKS:break
        cleaned.append(s)
    if not cleaned:
        raise HTTPException(422,"Унших боломжтой текст олдсонгүй. Зураг хэлбэртэй PDF-д OCR хэрэгтэй.")
    return cleaned

def replace_source(user_id,name,kind,passages,sheet_url=None,source_id=None):
    if not passages:raise HTTPException(422,"Мэдлэгийн текст олдсонгүй.")
    now=time.time()
    with core.db() as db:
        db.execute("BEGIN IMMEDIATE")
        if source_id:
            updated=db.execute("UPDATE reception_sources SET refreshed=?,last_error='' WHERE id=? AND user_id=? AND kind='sheet'",
                               (now,source_id,user_id)).rowcount
            if not updated:raise HTTPException(404,"Эх сурвалж олдсонгүй.")
            db.execute("DELETE FROM reception_passages WHERE source_id=?",(source_id,))
        else:
            source_id=core.uid()
            db.execute("INSERT INTO reception_sources(id,user_id,name,kind,sheet_url,created,refreshed) VALUES(?,?,?,?,?,?,?)",
                       (source_id,user_id,name[:200],kind,sheet_url,now,now))
        db.executemany("INSERT INTO reception_passages(source_id,line_no,text) VALUES(?,?,?)",
                       ((source_id,i,line) for i,line in enumerate(passages)))
    return {"id":source_id,"name":name[:200],"kind":kind,"status":"indexed","passages":len(passages)}

def list_sources(user_id):
    ensure_schema()
    with core.db() as db:
        rows=db.execute("""SELECT s.id,s.name,s.kind,s.refreshed,s.last_error,
                                (SELECT COUNT(*) FROM reception_passages p WHERE p.source_id=s.id) AS passages
                                FROM reception_sources s WHERE s.user_id=? ORDER BY s.created DESC LIMIT 100""",
                        (user_id,)).fetchall()
    return {"items":[dict(r) for r in rows]}

def remove_source(user_id,source_id):
    if not re.fullmatch(r"[a-f0-9]{32}",source_id):
        raise HTTPException(404,"Эх сурвалж олдсонгүй.")
    with core.db() as db:
        deleted=db.execute("DELETE FROM reception_sources WHERE user_id=? AND id=?",(user_id,source_id)).rowcount
        if not deleted:raise HTTPException(404,"Эх сурвалж олдсонгүй.")
    return {"ok":True}

def tokens(text):
    return {x for x in re.findall(r"[^\W_]+",str(text).lower(),re.UNICODE) if len(x)>=3 and x not in STOP}

def search(user_id,query,limit=4):
    if not isinstance(query,str) or not 2<=len(query.strip())<=350:
        raise HTTPException(422,"Хайх асуулт 2–350 тэмдэгттэй байх ёстой.")
    keys=tokens(query)
    if not keys:return {"matches":[],"message":"Тодорхой асуулт асууна уу."}
    with core.db() as db:
        rows=db.execute("""SELECT p.text,s.name,s.id FROM reception_passages p
                           JOIN reception_sources s ON s.id=p.source_id
                           WHERE s.user_id=? ORDER BY s.refreshed DESC LIMIT 10000""",(user_id,)).fetchall()
    scored=[]
    for r in rows:
        hit=keys&tokens(r["text"])
        if hit:
            score=len(hit)/len(keys)
            scored.append((score,len(hit),{"text":r["text"][:900],"source":r["name"],"source_id":r["id"]}))
    scored.sort(key=lambda x:(x[0],x[1]),reverse=True)
    return {"matches":[x[2] for x in scored[:limit]],"message":"Энэ бол эх сурвалжийн ишлэл. Хариуг нотолж чадахгүй бол ажилтнаас лавлана уу."}

def sheet_csv_url(url):
    if not isinstance(url,str) or len(url)>500 or "\\" in url or "@" in url:
        raise HTTPException(422,"Google Sheets холбоос буруу байна.")
    parts=urlsplit(url)
    if parts.scheme!="https" or parts.hostname!="docs.google.com" or parts.port not in (None,443) or parts.username or parts.password:
        raise HTTPException(422,"Зөвхөн https://docs.google.com/spreadsheets/d/... холбоос зөвшөөрнө.")
    match=SHEET_PATH.fullmatch(parts.path)
    if not match:raise HTTPException(422,"Google Sheets-ийн /spreadsheets/d/{id} хаяг оруулна уу.")
    g=parse_qs(parts.fragment).get("gid",parse_qs(parts.query).get("gid",["0"]))[0]
    if not re.fullmatch(r"[0-9]{1,12}",g):
        raise HTTPException(422,"Google Sheets gid буруу байна.")
    return "https://docs.google.com/spreadsheets/d/"+match.group(1)+"/export?format=csv&gid="+g

def sheet_account():
    """Operator-supplied service account (private JSON env var, never returned)."""
    raw=os.getenv("GOOGLE_SHEETS_SERVICE_ACCOUNT_JSON","").strip()
    if not raw:return None
    try:
        obj=json.loads(raw)
        if not isinstance(obj,dict) or obj.get("type")!="service_account" or not obj.get("client_email"):
            return None
        return obj
    except (ValueError,TypeError):
        return None

def google_status():
    account=sheet_account()
    return {"mode":"service_account" if account else "published_csv",
            "account_email":account["client_email"] if account else None,
            "sync_every_hours":6}

def private_sheet_csv(export_url,account):
    """Read privately shared Sheet with Google Sheets read-only API."""
    from google.oauth2 import service_account
    from google.auth.transport.requests import Request
    sheet_id=export_url.split("/d/")[1].split("/")[0]
    creds=service_account.Credentials.from_service_account_info(
        account,scopes=["https://www.googleapis.com/auth/spreadsheets.readonly"])
    creds.refresh(Request())
    api="https://sheets.googleapis.com/v4/spreadsheets/"+sheet_id+"/values/A1%3AZ2000"
    try:
        with httpx.Client(timeout=20,trust_env=False,follow_redirects=False) as client:
            response=client.get(api,headers={"Authorization":"Bearer "+creds.token})
            if response.status_code!=200:
                raise HTTPException(422,"Google Sheet-г service account имэйлд Viewer эрхээр хуваалцана уу.")
            values=response.json().get("values",[])
        out=io.StringIO()
        writer=csv.writer(out)
        for row in values[:2000]:writer.writerow([str(v)[:500] for v in row[:26]])
        raw=out.getvalue().encode("utf-8")
        if len(raw)>MAX_BYTES:raise HTTPException(413,"Google Sheets хэмжээ 5 MB-аас их.")
        return raw
    except HTTPException:raise
    except (ValueError,TypeError,httpx.HTTPError,Exception) as exc:
        raise HTTPException(502,"Google Sheets service account холболтыг шалгана уу.") from exc

def download_sheet(export_url):
    # Never allow arbitrary URLs or redirects: prevents user-controlled SSRF.
    if not export_url.startswith("https://docs.google.com/spreadsheets/d/") or sheet_csv_url(export_url)!=export_url:
        raise HTTPException(422,"Google Sheets холбоос зөвшөөрөгдөөгүй.")
    account=sheet_account()
    if account:
        return private_sheet_csv(export_url,account)
    try:
        with httpx.Client(timeout=20,follow_redirects=False,trust_env=False) as client:
            url=export_url
            # Google may redirect the CSV export to its own googleusercontent host.
            for hop in range(3):
                with client.stream("GET",url,headers={"Accept":"text/csv"}) as response:
                    if response.status_code in (301,302,303,307,308):
                        location=response.headers.get("location","")
                        next_url=urlsplit(httpx.URL(url).join(location).__str__())
                        host=next_url.hostname or ""
                        if next_url.scheme!="https" or not (host=="docs.google.com" or host=="googleusercontent.com" or host.endswith(".googleusercontent.com")):
                            raise HTTPException(422,"Google Sheets-ийн таталтын чиглүүлэлт зөвшөөрөгдөөгүй.")
                        url=next_url.geturl()
                        continue
                    if response.status_code!=200:
                        raise HTTPException(422,"Google Sheets CSV нийтлэх эрх болон файлын хандалтыг шалгана уу.")
                    raw=bytearray()
                    for chunk in response.iter_bytes():
                        raw.extend(chunk)
                        if len(raw)>MAX_BYTES:raise HTTPException(413,"Google Sheets 5 MB хязгаараас хэтэрлээ.")
                    return bytes(raw)
            raise HTTPException(422,"Google Sheets-ийн таталтын хаяг хэт олон удаа өөрчлөгдлөө.")
    except HTTPException:raise
    except httpx.HTTPError as exc:
        raise HTTPException(502,"Google Sheets холболт амжилтгүй. Нийтэлсэн CSV хаягийг шалгана уу.") from exc

def import_sheet(user_id,url):
    export=sheet_csv_url(url)
    data=download_sheet(export)
    passages=extract_document("sheet.csv",data)
    return replace_source(user_id,"Google Sheets: "+export.split("/d/")[1].split("/")[0][:45],"sheet",passages,export)

def interpret_time(text):
    from zoneinfo import ZoneInfo
    if not isinstance(text,str) or len(text)>500:
        raise HTTPException(422,"Хэлсэн цагийн мэдээлэл буруу байна.")
    now=datetime.now(ZoneInfo("Asia/Ulaanbaatar"))
    value=text.strip().lower()
    day=None
    explicit_date=re.search(r"\b(20\d{2}-\d{2}-\d{2})\b",value)
    if explicit_date:day=explicit_date.group(1)
    elif "нөгөөдөр" in value:day=(now+timedelta(days=2)).date().isoformat()
    elif "маргааш" in value:day=(now+timedelta(days=1)).date().isoformat()
    elif "өнөөдөр" in value:day=now.date().isoformat()
    time_found=None
    exact=re.search(r"(?<!\d)([01]?\d|2[0-3]):([0-5]\d)(?!\d)",value)
    if exact:
        hour=int(exact.group(1));minute=int(exact.group(2))
        if minute in (0,30):time_found=f"{hour:02d}:{minute:02d}"
    else:
        hour_match=re.search(r"(?<!\d)([01]?\d|2[0-3])\s*цаг",value)
        if hour_match:
            hour=int(hour_match.group(1))
            if hour==0 or hour>12:time_found=f"{hour:02d}:00"
            elif "өглөө" in value:time_found=f"{hour:02d}:00"
            elif "өдөр" in value or "орой" in value:
                time_found=f"{(hour+12 if hour<12 else hour):02d}:00"
            # "3 цаг" is unsafe (morning/afternoon): explicit confirmation required
    return {"date":day,"time":time_found,"timezone":"Asia/Ulaanbaatar",
            "confidence":"exact" if day and time_found else "clarify",
            "message":"Баталгаажуулж календарийн сул цаг шалгана уу." if day and time_found else "Өдөр болон өглөө/өдөр/орой, 24 цагийн цагийг тодруулна уу."}

async def periodic_sheet_sync():
    """Refresh published Sheets every six hours without prompting customer.
    Runs only in the Railway web service; stale content stays available on error.
    """
    while True:
        try:
            await asyncio.sleep(120)
            now=time.time()
            with core.db() as db:
                rows=db.execute("""SELECT id,user_id,sheet_url FROM reception_sources
                                   WHERE kind='sheet' AND refreshed<? ORDER BY refreshed LIMIT 10""",
                                (now-SHEET_INTERVAL,)).fetchall()
            for row in rows:
                try:
                    raw=await asyncio.to_thread(download_sheet,row["sheet_url"])
                    lines=await asyncio.to_thread(extract_document,"sheet.csv",raw)
                    replace_source(row["user_id"],"", "sheet",lines,source_id=row["id"])
                except Exception:
                    with core.db() as db:
                        db.execute("UPDATE reception_sources SET refreshed=?,last_error=? WHERE id=? AND user_id=?",
                                   (time.time(),"Синк хийхэд алдаа гарлаа; өмнөх мэдээллийг хадгалсан.",row["id"],row["user_id"]))
        except asyncio.CancelledError:raise
        except Exception:
            await asyncio.sleep(60)

def register_routes(app,admin_user,mutation_guard,json_body,throttle):
    @app.get("/api/reception/business")
    def get_business(request:Request):
        return read_business(admin_user(request)["user_id"])
    @app.post("/api/reception/business")
    async def set_business(request:Request):
        user=admin_user(request);mutation_guard(request,user)
        return save_business(user["user_id"],await json_body(request))
    @app.post("/api/reception/knowledge/upload")
    async def knowledge_upload(request:Request,file:UploadFile=File(...)):
        user=admin_user(request);mutation_guard(request,user)
        throttle("reception-file:"+user["user_id"],30,3600)
        name=Path(file.filename or "").name[:200]
        if Path(name).suffix.lower() not in EXTS:
            raise HTTPException(422,"Дэмжигдэх формат: TXT, CSV, XLSX, DOCX, PDF.")
        raw=await file.read(MAX_BYTES+1)
        if len(raw)>MAX_BYTES:
            raise HTTPException(413,"Файл хамгийн ихдээ 5 MB байна.")
        lines=await asyncio.to_thread(extract_document,name,raw)
        return replace_source(user["user_id"],name,"file",lines)
    @app.post("/api/reception/knowledge/sheet")
    async def knowledge_sheet(request:Request):
        user=admin_user(request);mutation_guard(request,user)
        throttle("reception-sheet:"+user["user_id"],10,3600)
        body=await json_body(request)
        return await asyncio.to_thread(import_sheet,user["user_id"],body.get("url"))
    @app.get("/api/reception/knowledge/google-status")
    def knowledge_google_status(request:Request):
        admin_user(request)
        return google_status()
    @app.get("/api/reception/knowledge/sources")
    def knowledge_sources(request:Request):
        return list_sources(admin_user(request)["user_id"])
    @app.delete("/api/reception/knowledge/sources/{source_id}")
    def knowledge_delete(source_id:str,request:Request):
        user=admin_user(request);mutation_guard(request,user)
        return remove_source(user["user_id"],source_id)
    @app.get("/api/reception/knowledge/search")
    def knowledge_search(request:Request,query:str):
        return search(admin_user(request)["user_id"],query)
    @app.post("/api/reception/time/interpret")
    async def knowledge_interpret(request:Request):
        user=admin_user(request);mutation_guard(request,user)
        return interpret_time((await json_body(request)).get("text"))
