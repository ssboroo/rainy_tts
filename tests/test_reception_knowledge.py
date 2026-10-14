"""Receptionist data ingestion, tenant boundary and business settings regression tests."""
import io
import os
import tempfile
import unittest
import time
from pathlib import Path
from unittest.mock import patch
from fastapi.testclient import TestClient
from app import core,server
from app import reception_knowledge as knowledge

class ReceptionKnowledgeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.old=core.DATA
        cls.temp=tempfile.TemporaryDirectory()
        core.DATA=Path(cls.temp.name)
        core.init()
        knowledge.ensure_schema()
        cls.client=TestClient(server.app)
    @classmethod
    def tearDownClass(cls):
        cls.client.close()
        core.DATA=cls.old
        cls.temp.cleanup()
    def setUp(self):
        self.client.cookies.clear()
        with server.RATE_LOCK:server.RATE.clear()
        self.email="kb-"+str(time.time_ns())+"@example.com"
        self.env=patch.dict(os.environ,{"ALLOW_REGISTRATION":"true","BILLING_ENABLED":"false","ADMIN_EMAILS":""})
        self.env.start()
        self.addCleanup(self.env.stop)
        server.ORIGIN="http://testserver"
        resp=self.client.post("/api/register",json={"email":self.email,"password":"strong-password-123"},headers={"Origin":"http://testserver"})
        self.assertEqual(resp.status_code,200,resp.text)
        self.headers={"Origin":"http://testserver","X-CSRF-Token":resp.json()["user"]["csrf"]}
        self.admin=patch.dict(os.environ,{"ADMIN_EMAILS":self.email})
        self.admin.start()
        self.addCleanup(self.admin.stop)
    def tearDown(self):
        with server.RATE_LOCK:server.RATE.clear()
    def post_file(self,content,name):
        return self.client.post("/api/reception/knowledge/upload",files={"file":(name,content)},headers=self.headers)
    def test_csv_utf8_import_and_grounded_answer_via_webhook(self):
        response=self.post_file("Асуулт,Хариулт\nАжиллах цаг,Бид Даваа Бямба 09:00-18:00 ажиллана\nҮйлчилгээ,Манай үйлчилгээний үнэ 50000 төгрөг".encode(),"faq.csv")
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(response.json()["status"],"indexed")
        config=self.client.get("/api/reception/config").json()
        from urllib.parse import urlsplit
        url=next(t["url"] for t in config["tools"] if t["name"]=="search_knowledge")
        self.client.cookies.clear()
        result=self.client.post(urlsplit(url).path,json={"query":"Ажиллах цаг хэд вэ"})
        self.assertEqual(result.status_code,200,result.text)
        self.assertTrue(result.json()["matches"],result.text)
        self.assertIn("09:00",result.json()["matches"][0]["text"])
        self.assertTrue(result.json()["matches"][0]["source"])
        missing=self.client.post(urlsplit(url).path,json={"query":"WiFi нууц үг"})
        self.assertEqual(missing.status_code,200)
        self.assertEqual(missing.json()["matches"],[])
    def test_business_profile_and_admin_boundary(self):
        r=self.client.post("/api/reception/business",json={"name":"Сарны салон","address":"УБ БЗД","contact_phone":"99112233","intro":"Өдөр бүр тавтай морилно уу"},headers=self.headers)
        self.assertEqual(r.status_code,200,r.text)
        got=self.client.get("/api/reception/business")
        self.assertEqual(got.json()["name"],"Сарны салон")
        with patch.dict(os.environ,{"ADMIN_EMAILS":"other@example.com"}):
            self.assertEqual(self.client.get("/api/reception/business").status_code,403)
            self.assertEqual(self.post_file(b"secret","secret.txt").status_code,403)
    def test_configurable_business_hours_affect_real_calendar(self):
        from datetime import datetime,timedelta
        from zoneinfo import ZoneInfo
        data={"name":"Нарны төв","address":"UB","contact_phone":"99112233",
              "intro":"Сайн байна уу","open_hour":11,"close_hour":16,"closed_weekdays":[6]}
        saved=self.client.post("/api/reception/business",json=data,headers=self.headers)
        self.assertEqual(saved.status_code,200,saved.text)
        self.assertEqual(saved.json()["open_hour"],11)
        from app import reception_booking
        services=self.client.get("/api/reception/services").json()["items"]
        day=datetime.now(ZoneInfo("Asia/Ulaanbaatar")).date()+timedelta(days=3)
        while day.weekday()==6:day+=timedelta(days=1)
        free=self.client.get("/api/reception/availability",params={"service_id":services[0]["id"],"date":day.isoformat()})
        self.assertEqual(free.status_code,200,free.text)
        self.assertIn("11:00",free.json()["slots"])
        self.assertNotIn("10:00",free.json()["slots"])
        self.assertNotIn("16:00",free.json()["slots"])
        invalid=self.client.post("/api/reception/appointments",json={
            "service_id":services[0]["id"],"date":day.isoformat(),"time":"10:00",
            "caller_name":"Бат","phone":"99112233","request_id":"hours-test-0001"
        },headers=self.headers)
        self.assertEqual(invalid.status_code,422,invalid.text)

    def test_docx_and_xlsx_import(self):
        from docx import Document
        from openpyxl import Workbook
        doc=Document();doc.add_paragraph("Манай төв шүдний үзлэг хийдэг.")
        data=io.BytesIO();doc.save(data)
        r=self.post_file(data.getvalue(),"info.docx")
        self.assertEqual(r.status_code,200,r.text)
        wb=Workbook();ws=wb.active;ws.append(["Асуулт","Хариулт"]);ws.append(["Хаяг","Монгол улс Улаанбаатар хот"])
        output=io.BytesIO();wb.save(output)
        r2=self.post_file(output.getvalue(),"info.xlsx")
        self.assertEqual(r2.status_code,200,r2.text)
        sources=self.client.get("/api/reception/knowledge/sources")
        self.assertEqual(sources.status_code,200)
        self.assertEqual(len(sources.json()["items"]),2)
    def test_pdf_import_and_limits(self):
        from pypdf import PdfWriter
        pdf=PdfWriter();pdf.add_blank_page(200,200)
        f=io.BytesIO();pdf.write(f)
        result=self.post_file(f.getvalue(),"empty.pdf")
        self.assertEqual(result.status_code,422,result.text)
        result=self.post_file(b"x"*30,"fake.pdf")
        self.assertEqual(result.status_code,422,result.text)
        result=self.post_file(b"x"*(6*1024*1024),"too-big.txt")
        self.assertEqual(result.status_code,413)
        self.assertEqual(self.post_file(b"abc","script.html").status_code,422)
    def test_google_sheet_url_allowlist_and_connector(self):
        bad=["http://localhost/private","https://docs.google.com.evil.org/spreadsheets/d/1234/edit",
             "https://docs.google.com/spreadsheets/d/abc123/../../secret",
             "https://evil.com/sheet.csv"]
        for url in bad:
            self.assertEqual(self.client.post("/api/reception/knowledge/sheet",json={"url":url},headers=self.headers).status_code,422,url)
        url="https://docs.google.com/spreadsheets/d/abcABC123456789_-/edit#gid=0"
        with patch.object(knowledge,"download_sheet",return_value=b"Q,A\nParking,We have covered parking"):
            result=self.client.post("/api/reception/knowledge/sheet",json={"url":url},headers=self.headers)
        self.assertEqual(result.status_code,200,result.text)
        self.assertEqual(result.json()["status"],"indexed")
        result2=self.client.get("/api/reception/knowledge/search",params={"query":"Parking"})
        self.assertEqual(result2.status_code,200)
        self.assertIn("covered parking",result2.json()["matches"][0]["text"])
    def test_disable_or_delete_source_removes_answer(self):
        uploaded=self.post_file(b"service,Wheel alignment offered","service.csv")
        self.assertEqual(uploaded.status_code,200,uploaded.text)
        sid=uploaded.json()["id"]
        results=self.client.get("/api/reception/knowledge/search",params={"query":"Wheel alignment"})
        self.assertTrue(results.json()["matches"])
        deleted=self.client.delete("/api/reception/knowledge/sources/"+sid,headers=self.headers)
        self.assertEqual(deleted.status_code,200,deleted.text)
        results=self.client.get("/api/reception/knowledge/search",params={"query":"Wheel alignment"})
        self.assertEqual(results.json()["matches"],[])
    def test_mongolian_day_and_time_interpretation_requires_disambiguation(self):
        parsed=self.client.post("/api/reception/time/interpret",json={"text":"маргааш 15:30 цагт"},headers=self.headers)
        self.assertEqual(parsed.status_code,200,parsed.text)
        self.assertEqual(parsed.json()["time"],"15:30")
        self.assertEqual(parsed.json()["confidence"],"exact")
        ambiguous=self.client.post("/api/reception/time/interpret",json={"text":"маргааш 3 цагт"},headers=self.headers)
        self.assertEqual(ambiguous.status_code,200)
        self.assertIsNone(ambiguous.json()["time"])
        self.assertEqual(ambiguous.json()["confidence"],"clarify")
